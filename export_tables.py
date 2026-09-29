#!/usr/bin/env python3
"""Export every analytical worksheet in the locked workbook as deterministic CSV."""

from __future__ import annotations

import csv
import hashlib
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
WORKBOOK = ROOT / "data" / "Article2_Consolidated_Tables.xlsx"
TABLE_ROOT = ROOT / "tables"

GROUPS = {
    "main": ("Table1_Sources", "Table2_Programs", "Table3_CPTAC"),
    "supplementary": (
        "S1_ProgramGenes",
        "S2_CPTAC_Flow",
        "S3_Primary",
        "S4a_FeatureSens",
        "S4b_Threshold1",
        "S4c_Pooled",
        "S4d_ResidualDiag",
        "S5_ResidualContext",
        "S6a_TCGA_Full",
        "S6b_TCGA_Summary",
        "S6c_TCGA_Anchors",
        "S6d_TCGA_Subtypes",
        "S7_APOLLO",
        "S8a_scRNA_Global",
        "S8b_scRNA_Source",
        "S9_Spatial_BRCA",
    ),
    "supporting": ("QA_Checks", "Derived_Heterogeneity", "Source_Alias"),
}


def cell_text(value: object) -> object:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return value


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def main() -> None:
    workbook = load_workbook(WORKBOOK, read_only=False, data_only=True)
    manifest: list[list[object]] = []

    for group, sheets in GROUPS.items():
        destination = TABLE_ROOT / group
        destination.mkdir(parents=True, exist_ok=True)
        for sheet_name in sheets:
            sheet = workbook[sheet_name]
            path = destination / f"{sheet_name}.csv"
            with path.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.writer(handle, lineterminator="\n")
                for row in sheet.iter_rows(
                    min_row=1,
                    max_row=sheet.max_row,
                    min_col=1,
                    max_col=sheet.max_column,
                    values_only=True,
                ):
                    writer.writerow([cell_text(value) for value in row])
            manifest.append(
                [
                    sheet_name,
                    group,
                    sheet.max_row - 1,
                    sheet.max_column,
                    path.relative_to(ROOT).as_posix(),
                    sha256(path),
                ]
            )

    manifest_path = TABLE_ROOT / "TABLE_EXPORT_MANIFEST.csv"
    with manifest_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["worksheet", "group", "records", "columns", "path", "sha256"])
        writer.writerows(manifest)

    print(f"Exported {len(manifest)} worksheets to {TABLE_ROOT}")
    print(manifest_path)


if __name__ == "__main__":
    main()
