#!/usr/bin/env python3
"""Validate the Article 002 technical release candidate and write its audit files."""

from __future__ import annotations

import csv
import hashlib
import json
import re
from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image


ROOT = Path(__file__).resolve().parent
WORKBOOK = ROOT / "data" / "Article2_Consolidated_Tables.xlsx"
VALIDATION_DIR = ROOT / "validation"
EXPECTED_SHEETS = [
    "Contents",
    "Table1_Sources",
    "Table2_Programs",
    "Table3_CPTAC",
    "S1_ProgramGenes",
    "S2_CPTAC_Flow",
    "S3_Primary",
    "S4a_FeatureSens",
    "S4b_Threshold1",
    "S4c_Pooled",
    "S4d_ResidualDiag",
    "S5_ResidualContext",
    "S6a_TCGA_Full",
    "S6b_TCGA_Summary",
    "S6c_TCGA_Anchors",
    "S6d_TCGA_Subtypes",
    "S7_APOLLO",
    "S8a_scRNA_Global",
    "S8b_scRNA_Source",
    "S9_Spatial_BRCA",
    "QA_Checks",
    "Derived_Heterogeneity",
    "Source_Alias",
]


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def rows_as_dicts(workbook, sheet_name: str) -> list[dict[str, object]]:
    sheet = workbook[sheet_name]
    header = [cell.value for cell in sheet[1]]
    return [dict(zip(header, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def workbook_cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def validate_csv_exports(workbook) -> int:
    manifest_path = ROOT / "tables" / "TABLE_EXPORT_MANIFEST.csv"
    with manifest_path.open(encoding="utf-8", newline="") as handle:
        entries = list(csv.DictReader(handle))
    require(len(entries) == 22, f"Expected 22 CSV exports, found {len(entries)}")
    for entry in entries:
        path = ROOT / entry["path"]
        require(path.is_file() and path.stat().st_size > 0, f"Missing CSV: {path}")
        require(sha256(path) == entry["sha256"], f"CSV checksum mismatch: {path}")
        with path.open(encoding="utf-8", newline="") as handle:
            csv_rows = list(csv.reader(handle))
        sheet = workbook[entry["worksheet"]]
        expected = [
            [workbook_cell_text(value) for value in row]
            for row in sheet.iter_rows(
                min_row=1,
                max_row=sheet.max_row,
                min_col=1,
                max_col=sheet.max_column,
                values_only=True,
            )
        ]
        require(csv_rows == expected, f"CSV values differ from workbook: {path}")
        require(int(entry["records"]) == sheet.max_row - 1, f"CSV row count mismatch: {path}")
        require(int(entry["columns"]) == sheet.max_column, f"CSV column count mismatch: {path}")
    return len(entries)


def validate_figures() -> dict[str, int]:
    counts: dict[str, int] = {}
    for directory, suffix, dpi in [
        ("PNG_300dpi", ".png", 300),
        ("SVG_editable", ".svg", None),
        ("TIFF_600dpi", ".tiff", 600),
    ]:
        files = sorted((ROOT / "generated" / directory).glob(f"*{suffix}"))
        require(len(files) == 10, f"Expected 10 {suffix} figures, found {len(files)}")
        for path in files:
            require(path.stat().st_size > 0, f"Empty figure: {path}")
            if dpi is not None:
                with Image.open(path) as image:
                    image.load()
                    require(round(image.info.get("dpi", (0, 0))[0]) == dpi, f"Unexpected DPI: {path}")
                    if suffix == ".tiff":
                        require(image.mode == "RGB", f"TIFF is not RGB: {path}")
        counts[directory] = len(files)
    return counts


def scan_portability_and_secrets() -> dict[str, int]:
    text_suffixes = {".py", ".md", ".txt", ".csv", ".json", ".cff", ".template", ".svg"}
    absolute_hits: list[str] = []
    secret_hits: list[str] = []
    absolute_pattern = re.compile(r"/" + r"Users/|/" + r"Volumes/|[A-Za-z]:\\\\")
    secret_pattern = re.compile(r"AKIA[0-9A-Z]{16}|ghp_[A-Za-z0-9]{30,}|sk-[A-Za-z0-9]{20,}|BEGIN (?:RSA |OPENSSH )?PRIVATE KEY")
    for path in ROOT.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in text_suffixes:
            continue
        if ".git" in path.parts:
            continue
        if path.name == "CHECKSUMS_SHA256.csv" or "workbook_renders" in path.parts:
            continue
        text = path.read_text(encoding="utf-8", errors="ignore")
        if absolute_pattern.search(text):
            absolute_hits.append(path.relative_to(ROOT).as_posix())
        if secret_pattern.search(text):
            secret_hits.append(path.relative_to(ROOT).as_posix())
    require(not absolute_hits, f"Absolute local paths found: {absolute_hits}")
    require(not secret_hits, f"Possible secrets found: {secret_hits}")
    return {"absolute_path_hits": 0, "secret_hits": 0}


def write_checksums() -> int:
    output = ROOT / "CHECKSUMS_SHA256.csv"
    files = [
        path
        for path in ROOT.rglob("*")
        if path.is_file()
        and path != output
        and ".git" not in path.parts
        and "workbook_renders" not in path.parts
        and "__pycache__" not in path.parts
    ]
    with output.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["sha256", "bytes", "path"])
        for path in sorted(files, key=lambda item: item.relative_to(ROOT).as_posix()):
            writer.writerow([sha256(path), path.stat().st_size, path.relative_to(ROOT).as_posix()])
    return len(files)


def main() -> None:
    workbook = load_workbook(WORKBOOK, read_only=False, data_only=True)
    require(workbook.sheetnames == EXPECTED_SHEETS, "Workbook sheet order or membership changed")

    primary = rows_as_dicts(workbook, "S3_Primary")
    require(len(primary) == 160, "Primary registry must contain 160 rows")
    evaluable = [row for row in primary if row["status"] == "evaluable"]
    require(len(evaluable) == 129, "Primary registry must contain 129 evaluable rows")
    layer_counts = Counter(row["comparison_layer"] for row in evaluable)
    require(layer_counts == Counter({"proteomics": 77, "phosphoproteomics": 52}), f"Unexpected layer counts: {layer_counts}")
    protein = [row for row in evaluable if row["comparison_layer"] == "proteomics"]
    phospho = [row for row in evaluable if row["comparison_layer"] == "phosphoproteomics"]
    require(sum(float(row["spearman_rho"]) > 0 for row in protein) == 77, "RNA–protein positivity mismatch")
    require(sum(float(row["q_value_bh"]) < 0.05 for row in protein) == 76, "RNA–protein q-value mismatch")
    require(sum(float(row["spearman_rho"]) > 0 for row in phospho) == 52, "RNA–phosphoprotein positivity mismatch")
    require(sum(float(row["q_value_bh"]) < 0.05 for row in phospho) == 52, "RNA–phosphoprotein q-value mismatch")

    tcga = rows_as_dicts(workbook, "S6a_TCGA_Full")
    require(len(tcga) == 5440, "TCGA registry must contain 5,440 rows")
    require(sum(float(row["q_value_bh"]) < 0.05 for row in tcga) == 4850, "TCGA q-value count mismatch")
    require(len(rows_as_dicts(workbook, "S6c_TCGA_Anchors")) == 17, "TCGA anchor count mismatch")

    apollo = rows_as_dicts(workbook, "S7_APOLLO")
    require(
        Counter(row["evidence_status"] for row in apollo)
        == Counter({"supported": 11, "null_result": 5, "opposite": 1, "not_evaluable": 7}),
        "APOLLO evidence-status count mismatch",
    )
    single_cell = rows_as_dicts(workbook, "S8b_scRNA_Source")
    require(len(single_cell) == 336, "Single-cell registry must contain 336 rows")
    require(sum(row["status"] == "evaluable" for row in single_cell) == 256, "Single-cell evaluable count mismatch")
    require(sum(bool(row["within_source_support_window"]) for row in single_cell) == 116, "Single-cell support count mismatch")
    spatial = rows_as_dicts(workbook, "S9_Spatial_BRCA")
    require(len(spatial) == 24 and {row["paired_sample_n"] for row in spatial} == {4}, "Spatial registry mismatch")

    csv_count = validate_csv_exports(workbook)
    figure_counts = validate_figures()
    scan_counts = scan_portability_and_secrets()

    manifest = json.loads((ROOT / "generated" / "run_manifest.json").read_text(encoding="utf-8"))
    require(manifest["workbook"] == "data/Article2_Consolidated_Tables.xlsx", "Run manifest is not portable")
    require(len(manifest["files"]) == 30 and all(not value.startswith("/") for value in manifest["files"]), "Run manifest output paths are not portable")
    require(manifest["package_versions"] == {"numpy": "2.5.1", "pandas": "3.0.5", "matplotlib": "3.11.1", "Pillow": "12.3.0"}, "Run environment differs from lock")

    manuscript = (ROOT / "manuscript" / "FULL_MANUSCRIPT_RELEASE_CANDIDATE.md").read_text(encoding="utf-8")
    require("All 77 evaluable RNA–protein correlations were positive, and 76 had q < 0.05" in manuscript, "Corrected RNA–protein claim missing")
    require("RNA–phosphoprotein/acetylation" not in manuscript, "Unsupported acetylation claim remains")
    require("prespecified" not in manuscript.lower() and "predefined" not in manuscript.lower(), "Unsupported prespecification wording remains")
    require("| Layer | Program | Evaluable cancers |" in manuscript, "Main tables were not populated")
    require("tables/supplementary/S6a_TCGA_Full.csv" in manuscript, "Supplementary index is incomplete")

    license_files = [ROOT / "LICENSE.md", ROOT / "LICENSE-CODE", ROOT / "LICENSE-CONTENT"]
    require(all(path.is_file() for path in license_files), "Scoped repository license files are incomplete")
    source_manifest = (ROOT / "provenance" / "SOURCE_MANIFEST_PUBLIC.csv").read_text(encoding="utf-8")
    require("NOT_ASSESSED" not in source_manifest, "Unassessed source redistribution rows remain")
    require(source_manifest.count("NOT_REDISTRIBUTED") == 105, "Source non-redistribution count mismatch")
    citation_path = ROOT / "CITATION.cff"
    require(citation_path.is_file(), "CITATION.cff is missing")
    citation = citation_path.read_text(encoding="utf-8")
    repository_url = "https://github.com/yananwei0-0/article002-immune-program-concordance"
    require(repository_url in citation, "Final repository URL is missing from CITATION.cff")

    remaining_blockers: list[str] = []
    if "Article 002 authors" in citation:
        remaining_blockers.append("review-anonymity decision and named author metadata")
    if not re.search(r"(?m)^(?:doi:|\s+- type: doi\s*$)", citation):
        remaining_blockers.append("archived release DOI")
    remaining_blockers.append("full final validation review and public-release authorization")
    public_release_ready = False

    VALIDATION_DIR.mkdir(parents=True, exist_ok=True)
    report = {
        "release": "1.0.0-rc1",
        "validated_on": date.today().isoformat(),
        "technical_validation": "PASS",
        "public_release_ready": public_release_ready,
        "resolved_release_metadata": [
            "scoped dual license",
            "source-terms and non-redistribution audit",
            "repository URL",
            "interim schema-valid CITATION.cff",
            "full technical validation rerun",
        ],
        "workbook_sheets": len(workbook.sheetnames),
        "data_sheets": len(workbook.sheetnames) - 1,
        "csv_exports": csv_count,
        "figure_counts": figure_counts,
        "scientific_invariants": {
            "cptac_registered": 160,
            "cptac_evaluable": 129,
            "rna_protein_positive": 77,
            "rna_protein_q_lt_0_05": 76,
            "rna_phosphoprotein_positive": 52,
            "rna_phosphoprotein_q_lt_0_05": 52,
            "tcga_registry": 5440,
            "tcga_q_lt_0_05": 4850,
            "single_cell_registry": 336,
            "spatial_registry": 24,
        },
        "portability_and_secret_scan": scan_counts,
        "remaining_blockers": remaining_blockers,
    }
    (VALIDATION_DIR / "VALIDATION.json").write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    audit = """# Article 002 release audit

Technical validation: **PASS on {validation_date}**
Public release ready: **NO — anonymity, named authors, and archived DOI remain**

Validated on {validation_date} against the locked consolidated workbook.

- 23 workbook sheets: one contents sheet and 22 analytical sheets.
- 22 CSV exports matched the workbook cell-for-cell and passed SHA-256 checks.
- Four main and six supplementary figures were present in PNG, SVG, and TIFF formats.
- PNG files passed the 300-dpi check; RGB TIFF files passed the 600-dpi check.
- The regenerated manifest used portable relative paths and the exact locked package versions.
- The manuscript claim was corrected to 77/77 positive RNA–protein correlations and 76/77 with q < 0.05.
- Unsupported acetylation and prespecification wording was removed.
- No absolute local paths or credential-like strings were detected in public text assets.
- The scoped dual license, repository URL, source-terms audit, and interim citation metadata were present.

The remaining blockers are listed in `PUBLIC_RELEASE_BLOCKERS.md` and do not require scientific recomputation.
""".format(validation_date=date.today().isoformat())
    (VALIDATION_DIR / "RELEASE_AUDIT.md").write_text(audit, encoding="utf-8")
    checksum_count = write_checksums()
    print(json.dumps({"technical_validation": "PASS", "csv_exports": csv_count, "checksummed_files": checksum_count}, ensure_ascii=False))


if __name__ == "__main__":
    main()
